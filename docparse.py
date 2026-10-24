"""Document parsing + local statistics for DocuLens.

Every parser returns a list of "units" (pages, slides, sheets, sections) so
answers can cite where information came from.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from html.parser import HTMLParser

SUPPORTED_TYPES = [
    "pdf", "docx", "pptx", "xlsx", "xls", "csv", "tsv",
    "txt", "md", "json", "html", "htm", "png", "jpg", "jpeg", "webp",
]
IMAGE_TYPES = {"png", "jpg", "jpeg", "webp"}
TABULAR_TYPES = {"csv", "tsv", "xlsx", "xls"}
MAX_TABLE_ROWS = 3000  # rows converted to text for AI (keeps prompts sane)


@dataclass
class ParsedDoc:
    name: str
    ext: str
    size: int
    digest: str
    units: list = field(default_factory=list)      # [(label, text)]
    tables: dict = field(default_factory=dict)     # {sheet: DataFrame} for preview
    warnings: list = field(default_factory=list)
    raw_bytes: bytes = b""
    needs_ocr: bool = False

    @property
    def full_text(self) -> str:
        return "\n\n".join(f"[{label}]\n{text}" for label, text in self.units if text.strip())

    @property
    def plain_text(self) -> str:
        return "\n\n".join(text for _, text in self.units if text.strip())

    @property
    def unit_name(self) -> str:
        return {"pdf": "Pages", "pptx": "Slides", "xlsx": "Sheets", "xls": "Sheets"}.get(self.ext, "Sections")


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _chunk_lines(text: str, per_unit: int = 6000) -> list:
    """Split plain text into labelled sections so citations stay meaningful."""
    text = text.strip()
    if len(text) <= per_unit:
        return [("Full text", text)] if text else []
    units, buf, n = [], "", 1
    for para in re.split(r"\n\s*\n", text):
        if len(buf) + len(para) > per_unit and buf:
            units.append((f"Section {n}", buf.strip()))
            buf, n = "", n + 1
        buf += para + "\n\n"
    if buf.strip():
        units.append((f"Section {n}", buf.strip()))
    return units


def _parse_pdf(doc: ParsedDoc):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(doc.raw_bytes))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            doc.warnings.append("This PDF is password-protected and could not be opened.")
            return
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        doc.units.append((f"Page {i}", text.strip()))
    if sum(len(t) for _, t in doc.units) < 50 * max(1, len(doc.units)) * 0.2:
        doc.needs_ocr = True
        doc.warnings.append(
            "Very little text was found. This looks like a scanned PDF, so AI text recognition (OCR) will be used."
        )


def _parse_docx(doc: ParsedDoc):
    from docx import Document

    d = Document(io.BytesIO(doc.raw_bytes))
    sections, current, title = [], [], "Introduction"
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            if current:
                sections.append((title, "\n".join(current)))
            title, current = txt[:70], []
        else:
            current.append(txt)
    if current:
        sections.append((title, "\n".join(current)))
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        sections.append((f"Table {ti}", "\n".join(rows)))
    doc.units = sections or []


def _parse_pptx(doc: ParsedDoc):
    from pptx import Presentation

    prs = Presentation(io.BytesIO(doc.raw_bytes))
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in r.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        doc.units.append((f"Slide {i}", "\n".join(parts)))


def _df_to_text(df, name: str) -> str:
    truncated = len(df) > MAX_TABLE_ROWS
    shown = df.head(MAX_TABLE_ROWS)
    out = [f"Columns: {', '.join(map(str, df.columns))}", f"Rows: {len(df)}"]
    try:
        desc = df.describe(include="all").T.fillna("").astype(str)
        out.append("Column summary:\n" + desc.to_string(max_colwidth=40))
    except Exception:
        pass
    out.append("Data:\n" + shown.to_csv(index=False))
    if truncated:
        out.append(f"(Only the first {MAX_TABLE_ROWS} of {len(df)} rows are included.)")
    return "\n".join(out)


def _parse_table(doc: ParsedDoc):
    import pandas as pd

    bio = io.BytesIO(doc.raw_bytes)
    if doc.ext in ("csv", "tsv"):
        sep = "\t" if doc.ext == "tsv" else None
        df = pd.read_csv(bio, sep=sep, engine="python", encoding_errors="replace")
        frames = {"Data": df}
    else:
        frames = pd.read_excel(bio, sheet_name=None)
    for sheet, df in frames.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        doc.tables[str(sheet)] = df
        doc.units.append((f"Sheet: {sheet}" if doc.ext in ("xlsx", "xls") else "Data", _df_to_text(df, sheet)))


def parse_document(name: str, data: bytes) -> ParsedDoc:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    doc = ParsedDoc(
        name=name, ext=ext, size=len(data),
        digest=hashlib.md5(data).hexdigest(), raw_bytes=data,
    )
    try:
        if ext == "pdf":
            _parse_pdf(doc)
        elif ext == "docx":
            _parse_docx(doc)
        elif ext == "pptx":
            _parse_pptx(doc)
        elif ext in TABULAR_TYPES:
            _parse_table(doc)
        elif ext in ("html", "htm"):
            p = _HTMLText()
            p.feed(_decode(data))
            doc.units = _chunk_lines(re.sub(r"\n{3,}", "\n\n", "".join(p.parts)))
        elif ext == "json":
            txt = _decode(data)
            try:
                txt = json.dumps(json.loads(txt), indent=2, ensure_ascii=False)
            except Exception:
                pass
            doc.units = _chunk_lines(txt)
        elif ext in IMAGE_TYPES:
            doc.needs_ocr = True
        elif ext in ("txt", "md"):
            doc.units = _chunk_lines(_decode(data))
        else:
            doc.warnings.append(f"Unsupported file type: .{ext}")
    except Exception as e:  # corrupted/odd files should not crash the app
        doc.warnings.append(f"Could not fully read this file: {e}")
    return doc


def apply_ocr_text(doc: ParsedDoc, text: str):
    doc.units = _chunk_lines(text) or []
    doc.needs_ocr = False


# ---------------------------------------------------------------- statistics
_STOP = set("""a about above after again against all also am an and any are as at be because been before being below
between both but by can could did do does doing down during each few for from further had has have having he her here
hers him his how i if in into is it its just me more most my no nor not now of off on once only or other our out over
own same she should so some such than that the their them then there these they this those through to too under until
up very was we were what when where which while who whom why will with would you your yours shall may might must per
via etc one two also use used using new""".split())


def _syllables(word: str) -> int:
    word = word.lower()
    groups = re.findall(r"[aeiouy]+", word)
    n = len(groups) - (1 if word.endswith("e") and len(groups) > 1 else 0)
    return max(1, n)


def compute_stats(doc: ParsedDoc) -> dict:
    text = doc.plain_text
    words = re.findall(r"\w+", text, flags=re.UNICODE)
    sentences = [s for s in re.split(r"(?<=[.!?\u0964])\s+", text) if len(s.split()) > 1]
    wc, sc = len(words), max(1, len(sentences))
    ascii_ratio = (sum(1 for c in text if ord(c) < 128) / len(text)) if text else 0
    freq = Counter(w.lower() for w in words if len(w) > 3 and not w.isdigit() and w.lower() not in _STOP)
    flesch = None
    if wc > 30 and ascii_ratio > 0.9:
        syl = sum(_syllables(w) for w in words)
        flesch = round(max(0, min(100, 206.835 - 1.015 * (wc / sc) - 84.6 * (syl / wc))), 1)
    return {
        "units": len([1 for _, t in doc.units if t.strip()]),
        "words": wc,
        "characters": len(text),
        "sentences": len(sentences),
        "read_minutes": max(1, round(wc / 220)) if wc else 0,
        "avg_sentence": round(wc / sc, 1) if wc else 0,
        "unique_ratio": round(len({w.lower() for w in words}) / wc * 100, 1) if wc else 0,
        "flesch": flesch,
        "keywords": freq.most_common(15),
    }


def readability_label(score):
    if score is None:
        return "n/a"
    if score >= 70:
        return "Easy"
    if score >= 50:
        return "Standard"
    if score >= 30:
        return "Difficult"
    return "Very difficult"


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"
