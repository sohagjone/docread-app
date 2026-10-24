"""Build downloadable reports (Markdown, HTML, Word) from one block structure."""
from __future__ import annotations

import html
import io
from datetime import datetime

from docparse import human_size, readability_label


def build_blocks(doc, stats: dict, analysis, chat: list | None = None) -> list:
    B = []
    title = (analysis.title if analysis and analysis.title else doc.name)
    B.append(("title", title))
    B.append(("meta", f"File: {doc.name}  |  Generated {datetime.now():%d %b %Y, %H:%M}  |  DocuLens AI"))

    B.append(("h2", "Document at a glance"))
    rows = [
        ["File size", human_size(doc.size)],
        [doc.unit_name, str(stats["units"])],
        ["Words", f"{stats['words']:,}"],
        ["Estimated reading time", f"{stats['read_minutes']} min"],
        ["Average sentence length", f"{stats['avg_sentence']} words"],
    ]
    if stats.get("flesch") is not None:
        rows.append(["Readability (Flesch)", f"{stats['flesch']} ({readability_label(stats['flesch'])})"])
    if analysis:
        if analysis.document_type:
            rows.insert(0, ["Document type", analysis.document_type])
        if analysis.language:
            rows.append(["Language", analysis.language])
        if analysis.clarity_score:
            rows.append(["Clarity score", f"{analysis.clarity_score}/10"])
    B.append(("table", ["Metric", "Value"], rows))
    if stats["keywords"]:
        B.append(("p", "Top keywords: " + ", ".join(f"{w} ({c})" for w, c in stats["keywords"][:10])))

    if analysis:
        a = analysis
        if a.purpose:
            B.append(("h2", "Purpose")); B.append(("p", a.purpose))
        if a.intended_audience:
            B.append(("p", f"Intended audience: {a.intended_audience}"))
        if a.executive_summary:
            B.append(("h2", "Executive summary"))
            for para in [p.strip() for p in a.executive_summary.split("\n") if p.strip()]:
                B.append(("p", para))
        if a.key_points:
            B.append(("h2", "Key points")); B.append(("bullets", a.key_points))
        if a.topics:
            B.append(("h2", "Main topics"))
            B.append(("table", ["Topic", "Description"], [[t.name, t.description] for t in a.topics]))
        if a.key_facts:
            B.append(("h2", "Key facts and figures"))
            B.append(("table", ["Item", "Value", "Context"], [[f.label, f.value, f.context] for f in a.key_facts]))
        if a.entities:
            B.append(("h2", "People, organisations and places"))
            B.append(("table", ["Name", "Type", "Role in document"], [[e.name, e.type, e.context] for e in a.entities]))
        if a.tone or a.sentiment_explanation:
            B.append(("h2", "Tone and sentiment"))
            B.append(("p", f"Tone: {a.tone}. Overall sentiment: {a.sentiment}. {a.sentiment_explanation}"))
        if a.risks:
            B.append(("h2", "Risks and concerns"))
            B.append(("table", ["Severity", "Issue", "Explanation"], [[r.severity, r.issue, r.explanation] for r in a.risks]))
        if a.action_items:
            B.append(("h2", "Action items"))
            B.append(("table", ["Action", "Owner", "Deadline"], [[x.action, x.owner or "-", x.deadline or "-"] for x in a.action_items]))
        if a.strengths:
            B.append(("h2", "Strengths")); B.append(("bullets", a.strengths))
        if a.gaps_or_weaknesses:
            B.append(("h2", "Gaps and weaknesses")); B.append(("bullets", a.gaps_or_weaknesses))
        if a.recommendations:
            B.append(("h2", "Recommendations")); B.append(("bullets", a.recommendations))
        if a.clarity_notes:
            B.append(("h2", "Clarity and organisation")); B.append(("p", a.clarity_notes))
        if a.glossary:
            B.append(("h2", "Glossary"))
            B.append(("table", ["Term", "Definition"], [[t.term, t.definition] for t in a.glossary]))
    if chat:
        B.append(("h2", "Questions and answers"))
        for m in chat:
            B.append(("qa", m["role"], m["content"]))
    return B


# ------------------------------------------------------------------ Markdown
def to_markdown(blocks: list) -> str:
    out = []
    for b in blocks:
        k = b[0]
        if k == "title":
            out.append(f"# {b[1]}\n")
        elif k == "meta":
            out.append(f"_{b[1]}_\n")
        elif k == "h2":
            out.append(f"\n## {b[1]}\n")
        elif k == "p":
            out.append(b[1] + "\n")
        elif k == "bullets":
            out.append("\n".join(f"- {x}" for x in b[1]) + "\n")
        elif k == "table":
            esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")
            out.append("| " + " | ".join(b[1]) + " |")
            out.append("|" + "---|" * len(b[1]))
            out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in b[2]]
            out.append("")
        elif k == "qa":
            out.append(f"**{'You' if b[1] == 'user' else 'DocuLens'}:** {b[2]}\n")
    return "\n".join(out)


# ---------------------------------------------------------------------- HTML
_CSS = """
body{font-family:Inter,Segoe UI,Arial,sans-serif;color:#1e293b;max-width:880px;margin:40px auto;padding:0 24px;line-height:1.65}
h1{color:#0b2545;border-bottom:4px solid #0e9aa7;padding-bottom:10px;margin-bottom:4px}
h2{color:#0b2545;margin-top:34px;border-left:5px solid #0e9aa7;padding-left:12px}
.meta{color:#64748b;font-size:13px;margin-bottom:24px}
table{border-collapse:collapse;width:100%;margin:12px 0;font-size:14px}
th{background:#0b2545;color:#fff;text-align:left;padding:9px 12px}
td{border-bottom:1px solid #e2e8f0;padding:8px 12px;vertical-align:top}
tr:nth-child(even) td{background:#f8fafc}
.q{background:#eff6ff;border-radius:10px;padding:10px 14px;margin:10px 0}
.a{background:#f1f5f9;border-radius:10px;padding:10px 14px;margin:10px 0}
footer{margin-top:40px;color:#94a3b8;font-size:12px;text-align:center}
"""


def to_html(blocks: list) -> str:
    e = lambda s: html.escape(str(s)).replace("\n", "<br>")
    out = []
    for b in blocks:
        k = b[0]
        if k == "title":
            out.append(f"<h1>{e(b[1])}</h1>")
        elif k == "meta":
            out.append(f"<div class='meta'>{e(b[1])}</div>")
        elif k == "h2":
            out.append(f"<h2>{e(b[1])}</h2>")
        elif k == "p":
            out.append(f"<p>{e(b[1])}</p>")
        elif k == "bullets":
            out.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in b[1]) + "</ul>")
        elif k == "table":
            head = "".join(f"<th>{e(h)}</th>" for h in b[1])
            body = "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in b[2])
            out.append(f"<table><tr>{head}</tr>{body}</table>")
        elif k == "qa":
            cls, who = ("q", "You") if b[1] == "user" else ("a", "DocuLens")
            out.append(f"<div class='{cls}'><b>{who}:</b> {e(b[2])}</div>")
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>Document report</title>"
        f"<style>{_CSS}</style></head><body>{''.join(out)}<footer>Generated by DocuLens AI</footer></body></html>"
    )


# ---------------------------------------------------------------------- Word
def to_docx(blocks: list) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    NAVY, TEAL = RGBColor(0x0B, 0x25, 0x45), RGBColor(0x0E, 0x9A, 0xA7)
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    for b in blocks:
        k = b[0]
        if k == "title":
            p = d.add_paragraph(); r = p.add_run(b[1]); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
        elif k == "meta":
            p = d.add_paragraph(); r = p.add_run(b[1]); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        elif k == "h2":
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(14)
            r = p.add_run(b[1]); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = TEAL
        elif k == "p":
            d.add_paragraph(b[1])
        elif k == "bullets":
            for x in b[1]:
                d.add_paragraph(str(x), style="List Bullet")
        elif k == "table":
            t = d.add_table(rows=1, cols=len(b[1])); t.style = "Table Grid"
            for i, h in enumerate(b[1]):
                c = t.rows[0].cells[i]; c.text = ""
                run = c.paragraphs[0].add_run(h); run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
                shade(c, "0B2545")
            for row in b[2]:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
            d.add_paragraph()
        elif k == "qa":
            p = d.add_paragraph(); r = p.add_run(("You: " if b[1] == "user" else "DocuLens: ")); r.bold = True
            p.add_run(b[2])
    foot = d.add_paragraph("Generated by DocuLens AI"); foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()
